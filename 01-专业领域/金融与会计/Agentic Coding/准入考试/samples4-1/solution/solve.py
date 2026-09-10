from copy import copy
import os
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


SHEET_DIR = Path(os.environ.get("SHEET_DIR", "/app/sheet"))
START_ROW = 6
START_COLUMN = 5  # E
END_COLUMN = 28  # AB


def output_path(input_path: Path) -> Path:
    return input_path.with_name(f"{input_path.stem}_result{input_path.suffix}")


def build_formula(column: int, row: int) -> str:
    column_letter = get_column_letter(column)
    return (
        f"=(MOD(YEARFRAC({column_letter}$4,EOMONTH($D{row},0))*12,12/$C{row})=0)"
        f"*(EDATE($D{row},1)>{column_letter}$4)*$B{row}"
        f"+(EOMONTH($D{row},0)={column_letter}$4)*$A{row}"
    )


def process_workbook(input_path: Path) -> None:
    workbook = load_workbook(input_path)
    worksheet = workbook.worksheets[0]

    last_row = max(START_ROW - 1, worksheet.max_row)
    for row in range(START_ROW, last_row + 1):
        if all(worksheet.cell(row=row, column=column).value is None for column in range(1, 5)):
            continue

        for column in range(START_COLUMN, END_COLUMN + 1):
            cell = worksheet.cell(row=row, column=column)
            cell.value = build_formula(column, row)

            alignment = copy(cell.alignment)
            alignment.horizontal = "center"
            cell.alignment = alignment

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(output_path(input_path))


def main() -> None:
    inputs = sorted(SHEET_DIR.glob("*_init.xlsx"))
    if not inputs:
        raise FileNotFoundError(f"No *_init.xlsx files found in {SHEET_DIR}")

    for input_path in inputs:
        process_workbook(input_path)


if __name__ == "__main__":
    main()
